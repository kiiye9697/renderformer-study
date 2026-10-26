from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
import re
root=Path(__file__).parent
equation_keys={v:k for k,v in json.loads((root/'equations.json').read_text(encoding='utf-8')).items()}
import sys
private_mode='--private' in sys.argv
profile=json.loads((root/'private'/'profile.json').read_text(encoding='utf-8-sig')) if private_mode else {}
out_dir=root/'private' if private_mode else root
output_stem='RenderFormer-期末报告' if '--final' in sys.argv else 'RenderFormer-期末报告-工作稿'
def cover_text(value):
    replacements={'班    级：________________':'班    级：'+profile.get('class','________________'),'学生姓名：________________':'学生姓名：'+profile.get('name','________________'),'学    号：________________':'学    号：'+profile.get('student_id','________________'),'完成日期：2026 年 ____ 月 ____ 日':'完成日期：'+profile.get('completion_date','2026 年 ____ 月 ____ 日')}
    return replacements.get(value,value)
text=(root/'report-draft.md').read_text(encoding='utf-8-sig')
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.3); sec.bottom_margin=Cm(2.3)
sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name]
    st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
    st.font.size=Pt(12)
    st.paragraph_format.line_spacing=1
    st.paragraph_format.space_after=Pt(0)
    st.paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.bold=False
    st.paragraph_format.space_before=Pt(6)
    spacing=st.element.get_or_add_pPr().find(qn('w:spacing'))
    spacing.set(qn('w:beforeLines'),'50')
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(.85)
# Text-derived cover reconstruction; original binding-box geometry awaits visual validation.
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph('数据科学期末报告','Title'); p.alignment=1; p.runs[0].font.size=Pt(24)
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph('Transformer 学习方法及其在三角网格神经渲染中的应用','Title'); p.alignment=1; p.runs[0].font.size=Pt(18)
p=doc.add_paragraph('以 RenderFormer 为例','Title'); p.alignment=1; p.runs[0].font.size=Pt(16)
for _ in range(3): doc.add_paragraph()
for t in ['院    系：计算机科学与技术','班    级：________________','学生姓名：________________','学    号：________________','完成日期：2026 年 ____ 月 ____ 日']:
    p=doc.add_paragraph(cover_text(t)); p.alignment=1; p.paragraph_format.space_after=Pt(14)
for _ in range(2):doc.add_paragraph()
p=doc.add_paragraph('大连理工大学计算机科学与技术学院'); p.alignment=1
# Seven body pages intended; actual pagination must be checked by renderer.
page_break_titles={'二 Transformer 的表示与注意力计算','三 RenderFormer 的两阶段网络','四 数据采集与监督学习过程','五 复现实验设计与数据流程','六 评价指标与可复现记录','七 实验结果与讨论'}
doc.add_page_break()
for line in text.splitlines()[3:]:
    if not line.strip():continue
    if line=='<!--results-table-->':
        rows=json.loads((root/'result-table.json').read_text(encoding='utf-8'))
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
        for vals in rows:
            cells=table.add_row().cells
            for cell,value in zip(cells,vals):
                cell.text=str(value)
                for pp in cell.paragraphs:
                    pp.paragraph_format.first_line_indent=Pt(0)
                    for rr in pp.runs:rr.font.size=Pt(10)
        doc.add_paragraph()
    elif line.startswith('!['):
        m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
        if not m:raise ValueError('Invalid figure marker')
        p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(root/m.group(2)),width=Cm(15.5))
        cp=doc.add_paragraph(m.group(1));cp.alignment=1;cp.paragraph_format.first_line_indent=Pt(0)
        for rr in cp.runs:rr.font.size=Pt(10)
    elif line.startswith('$$'):
        key=equation_keys[line.strip('$')]
        p=doc.add_paragraph();p.alignment=1
        p.add_run().add_picture(str(root/'equations'/(key+'.png')),width=Cm(13))
    elif line.startswith('## '):
        h=line[3:]
        # Natural flow after the cover; page count must be rendered.
        doc.add_paragraph(h,'Heading 1')
    elif not line.startswith('#'):
        p=doc.add_paragraph(line)
        if line.startswith('表 1'):p.paragraph_format.keep_with_next=True
        if line.startswith('关键词') or line.startswith('['):p.paragraph_format.first_line_indent=Pt(0)
# Footer number intentionally normal page field, including cover, pending template fidelity check.
f=sec.footer.paragraphs[0]; f.alignment=1
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
doc.core_properties.title='Transformer 学习方法及其在三角网格神经渲染中的应用'
doc.core_properties.author=''
doc.core_properties.comments=''
out=out_dir/(output_stem+'.docx');doc.save(out)
checks={'output':str(out),'body_character_count':len(text),'paragraph_count':len(doc.paragraphs),'pagination':'natural flow after cover','actual_page_count':None,'render_verified':False,'template_visual_fidelity_verified':False,'experimental_results_included':'<!--results-table-->' in text,'pending':['原封面视觉精确匹配未核验','Word逐页渲染未完成']}
(out_dir/'docx-structural-check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False))
